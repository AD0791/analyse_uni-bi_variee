from docx import Document
from docx.shared import Pt
import math  # For proper log10 calculation

# Créer un nouveau document Word
doc = Document()

# Ajouter un titre principal
title = doc.add_heading('Calcul des classes avec la formule de Sturges', level=1)
title.alignment = 1  # Centrer le titre

# Ajouter un paragraphe d'introduction
doc.add_paragraph(
    "La formule de Sturges est utilisée pour déterminer le nombre optimal de classes (ou intervalles) dans un histogramme. "
    "Elle permet d'organiser les données en groupes homogènes et de faciliter leur interprétation. Voici une explication détaillée de la méthode et les calculs effectués pour chaque variable quantitative."
)

# Ajouter un sous-titre pour la formule de Sturges
doc.add_heading('Formule de Sturges', level=2)

# Ajouter la formule - utiliser des strings bruts (r) pour éviter les problèmes d'échappement
doc.add_paragraph(
    r"La formule de Sturges est donnée par :"
)
doc.add_paragraph(
    r"k = 1 + 3.322 · log₁₀(n)"
)
doc.add_paragraph(
    r"Où :"
)
doc.add_paragraph(
    r"- k : Nombre de classes."
)
doc.add_paragraph(
    r"- n : Nombre total d'observations dans le jeu de données."
)
doc.add_paragraph(
    r"- log₁₀ : Logarithme en base 10."
)

# Ajouter un sous-titre pour la méthode de calcul des classes
doc.add_heading('Méthode de calcul des classes', level=2)

# Ajouter les étapes de calcul
doc.add_paragraph(
    r"Une fois le nombre de classes k déterminé, on calcule l'étendue des données et la largeur des classes comme suit :"
)
doc.add_paragraph(
    r"1. Étendue des données :"
)
doc.add_paragraph(
    r"Étendue = Valeur maximale - Valeur minimale"
)
doc.add_paragraph(
    r"2. Largeur des classes :"
)
doc.add_paragraph(
    r"Largeur = Étendue / k"
)
doc.add_paragraph(
    r"3. Limites des classes :"
)
doc.add_paragraph(
    r"Les limites des classes sont déterminées en ajoutant la largeur à la valeur minimale, puis en répétant cette opération pour chaque classe."
)

# Ajouter un sous-titre pour les calculs effectués
doc.add_heading('Calculs effectués pour chaque variable', level=2)

# Fonction pour calculer les classes avec la formule de Sturges
def calcul_classes_sturges(variable, nom_variable, valeur_min, valeur_max, n_observations):
    # Calcul du nombre de classes (k) - utiliser math.log10 pour un calcul correct
    k = 1 + 3.322 * math.log10(n_observations)
    k = round(k)  # Arrondir à l'entier le plus proche

    # Calcul de l'étendue
    etendue = valeur_max - valeur_min

    # Calcul de la largeur des classes
    largeur = etendue / k

    # Détermination des limites des classes
    limites_classes = []
    limite_inf = valeur_min
    for i in range(k):
        limite_sup = limite_inf + largeur
        limites_classes.append(f"[{limite_inf:.2f}, {limite_sup:.2f}[")
        limite_inf = limite_sup

    # Ajouter les résultats au document
    doc.add_heading(f'Variable : {nom_variable}', level=3)
    
    p = doc.add_paragraph()
    p.add_run(f"1. Nombre d'observations (n) : {n_observations}\n")
    p.add_run(f"2. Nombre de classes (k) : {k}\n")
    p.add_run(f"3. Étendue des données : {etendue:.2f}\n")
    p.add_run(f"4. Largeur des classes : {largeur:.2f}\n")
    p.add_run("5. Limites des classes :\n")
    
    for limite in limites_classes:
        doc.add_paragraph(limite)

# Données pour chaque variable
variables = [
    {"nom": "Âge (Q2)", "min": 27, "max": 73, "n": 41},
    {"nom": "Nombre de personnes sous le même toit (Q5)", "min": 2, "max": 8, "n": 41},
    {"nom": "Montant du crédit (Q10)", "min": 50000, "max": 150000, "n": 41},
    {"nom": "Temps d'approbation du crédit (Q11)", "min": 30, "max": 75, "n": 41},
]

# Calcul des classes pour chaque variable
for var in variables:
    calcul_classes_sturges(var["nom"], var["nom"], var["min"], var["max"], var["n"])

# Enregistrer le document
doc.save('Calcul_Classes_Sturges.docx')
