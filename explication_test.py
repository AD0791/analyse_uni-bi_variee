from docx import Document
from docx.shared import Inches

# Créer un nouveau document Word
doc = Document()

# Ajouter un titre principal
title = doc.add_heading('Explication des tests statistiques : Khi-deux et ANOVA', level=1)
title.alignment = 1  # Centrer le titre

# Ajouter un paragraphe d'introduction
doc.add_paragraph(
    "Ce document explique les tests statistiques du Khi-deux et de l'ANOVA, leur utilisation, et leur interprétation dans le cadre de l'étude sur l'impact des crédits agricoles sur les agriculteurs de Thomassique. "
    "Il précise également les variables à considérer et les tests à appliquer en fonction des objectifs du devoir."
)

# Ajouter un sous-titre pour le test du Khi-deux
doc.add_heading('1. Test du Khi-deux (χ²)', level=2)

# Explication du test du Khi-deux
doc.add_paragraph(
    "Le test du Khi-deux est utilisé pour analyser la relation entre deux variables catégorielles. Il permet de déterminer s'il existe une association significative entre ces variables."
)

doc.add_paragraph(
    "Quand utiliser le test du Khi-deux ?\n"
    "- Lorsque vous avez deux variables catégorielles.\n"
    "- Par exemple :\n"
    "  - Variable 1 : Accès au crédit agricole (Oui/Non).\n"
    "  - Variable 2 : Amélioration des revenus agricoles (Oui/Non)."
)

doc.add_paragraph(
    "Hypothèses du test du Khi-deux :\n"
    "- Hypothèse nulle (H₀) : Il n'y a pas d'association entre les deux variables.\n"
    "- Hypothèse alternative (H₁) : Il existe une association entre les deux variables."
)

doc.add_paragraph(
    "Interprétation des résultats :\n"
    "- Si la p-value est inférieure à 0.05, on rejette l'hypothèse nulle et on conclut qu'il existe une association significative entre les variables.\n"
    "- Si la p-value est supérieure à 0.05, on ne peut pas rejeter l'hypothèse nulle, et on conclut qu'il n'y a pas d'association significative."
)

# Exemple concret dans votre étude
doc.add_heading('Exemple concret dans votre étude :', level=3)
doc.add_paragraph(
    "Question : Les agriculteurs qui ont accès au crédit agricole sont-ils plus susceptibles de voir leurs revenus augmenter que ceux qui n'y ont pas accès ?\n"
    "Variables :\n"
    "- Variable 1 : Accès au crédit (Oui/Non).\n"
    "- Variable 2 : Amélioration des revenus (Oui/Non).\n"
    "Test : Test du Khi-deux."
)

# Ajouter un sous-titre pour l'ANOVA
doc.add_heading('2. ANOVA (Analyse de Variance)', level=2)

# Explication de l'ANOVA
doc.add_paragraph(
    "L'ANOVA est utilisée pour comparer les moyennes d'une variable quantitative entre trois groupes ou plus. Elle permet de déterminer s'il existe des différences significatives entre les moyennes des groupes."
)

doc.add_paragraph(
    "Quand utiliser l'ANOVA ?\n"
    "- Lorsque vous avez une variable quantitative et une variable catégorielle avec trois groupes ou plus.\n"
    "- Par exemple :\n"
    "  - Variable quantitative : Montant du crédit agricole.\n"
    "  - Groupes : Types de cultures (Légumineuses, Maraîchage, Riziculture, etc.)."
)

doc.add_paragraph(
    "Hypothèses de l'ANOVA :\n"
    "- Hypothèse nulle (H₀) : Les moyennes des groupes sont égales.\n"
    "- Hypothèse alternative (H₁) : Au moins une des moyennes des groupes est différente."
)

doc.add_paragraph(
    "Interprétation des résultats :\n"
    "- Si la p-value est inférieure à 0.05, on rejette l'hypothèse nulle et on conclut qu'il existe des différences significatives entre les moyennes des groupes.\n"
    "- Si la p-value est supérieure à 0.05, on ne peut pas rejeter l'hypothèse nulle, et on conclut qu'il n'y a pas de différences significatives."
)

# Exemple concret dans votre étude
doc.add_heading('Exemple concret dans votre étude :', level=3)
doc.add_paragraph(
    "Question : Le montant moyen du crédit agricole diffère-t-il selon le type de culture pratiquée ?\n"
    "Variables :\n"
    "- Variable quantitative : Montant du crédit.\n"
    "- Groupes : Types de cultures (Légumineuses, Maraîchage, etc.).\n"
    "Test : ANOVA."
)

# Ajouter un sous-titre pour les variables et tests à considérer
doc.add_heading('3. Variables et tests à considérer dans votre étude', level=2)

# Explication des variables et tests
doc.add_paragraph(
    "Voici une liste des variables à considérer et des tests à appliquer en fonction des objectifs de votre étude :"
)

doc.add_paragraph(
    "1. **Accès au crédit agricole et amélioration des revenus**\n"
    "- Variables :\n"
    "  - Variable 1 : Accès au crédit (Oui/Non).\n"
    "  - Variable 2 : Amélioration des revenus (Oui/Non).\n"
    "- Test : Test du Khi-deux.\n"
    "- Objectif : Déterminer si l'accès au crédit est associé à une amélioration des revenus."
)

doc.add_paragraph(
    "2. **Appartenance à une coopérative agricole et sécurité alimentaire**\n"
    "- Variables :\n"
    "  - Variable 1 : Appartenance à une coopérative (Oui/Non).\n"
    "  - Variable 2 : Amélioration de la sécurité alimentaire (Oui/Non).\n"
    "- Test : Test du Khi-deux.\n"
    "- Objectif : Déterminer si l'appartenance à une coopérative est associée à une amélioration de la sécurité alimentaire."
)

doc.add_paragraph(
    "3. **Montant du crédit agricole selon le type de culture**\n"
    "- Variables :\n"
    "  - Variable quantitative : Montant du crédit.\n"
    "  - Groupes : Types de cultures (Légumineuses, Maraîchage, etc.).\n"
    "- Test : ANOVA.\n"
    "- Objectif : Déterminer si le montant moyen du crédit diffère selon le type de culture."
)

doc.add_paragraph(
    "4. **Temps d'approbation du crédit selon la commune**\n"
    "- Variables :\n"
    "  - Variable quantitative : Temps d'approbation du crédit.\n"
    "  - Groupes : Communes (Thomassique, etc.).\n"
    "- Test : ANOVA.\n"
    "- Objectif : Déterminer si le temps moyen d'approbation du crédit diffère selon la commune."
)

# Ajouter un sous-titre pour l'interprétation des p-values
doc.add_heading('4. Interprétation des p-values', level=2)

# Explication de l'interprétation des p-values
doc.add_paragraph(
    "La p-value est une mesure qui permet de déterminer si les résultats observés sont statistiquement significatifs. Voici comment interpréter les p-values :"
)

doc.add_paragraph(
    "- **Si la p-value est inférieure à 0.05** :\n"
    "  - On rejette l'hypothèse nulle (H₀).\n"
    "  - On conclut qu'il existe une association significative (Khi-deux) ou une différence significative (ANOVA)."
)

doc.add_paragraph(
    "- **Si la p-value est supérieure à 0.05** :\n"
    "  - On ne peut pas rejeter l'hypothèse nulle (H₀).\n"
    "  - On conclut qu'il n'y a pas d'association significative (Khi-deux) ou pas de différence significative (ANOVA)."
)

# Enregistrer le document
doc.save('Explication_Tests_Statistiques.docx')
