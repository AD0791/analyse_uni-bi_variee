from docx import Document
from docx.shared import Pt
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

# Créer un nouveau document Word
doc = Document()

# Ajouter un titre principal
title = doc.add_heading('Étape 1 : Sélection des variables pertinentes', level=1)
title.alignment = 1  # Centrer le titre

# Ajouter un paragraphe d'introduction
doc.add_paragraph(
    "Nous devons sélectionner au moins 12 variables pertinentes pour l'objet d'étude, qui est l'impact des crédits agricoles sur les agriculteurs de Thomassique. "
    "Voici une proposition de variables basée sur le questionnaire et les indicateurs mentionnés dans le document Word."
)

# Ajouter un sous-titre pour le tableau
doc.add_heading('Tableau des variables sélectionnées', level=2)

# Créer un tableau
table = doc.add_table(rows=1, cols=3)
table.style = 'Table Grid'  # Ajouter des bordures au tableau

# Ajouter les en-têtes du tableau
header_cells = table.rows[0].cells
header_cells[0].text = 'Nom des variables'
header_cells[1].text = 'Type de données/modalités'
header_cells[2].text = 'Commentaires/explication'

# Données du tableau
variables = [
    ("Q2-Veuillez indiquer votre Age.", "Quantitative (continue)", "Âge des répondants."),
    ("Q3-Quel est votre sexe?", "Qualitative (catégorielle)", "Sexe des répondants (Homme/Femme)."),
    ("Q5-Combien de personnes vivent sous le même toit?", "Quantitative (discrète)", "Taille du ménage."),
    ("Q8-Avez-vous déjà contracté un crédit agricol?", "Qualitative (catégorielle)", "Accès au crédit (Oui/Non)."),
    ("Q9-Si oui, combien de fois avez-vous obtenu un crédit agricole ces 3 dernières années?", "Qualitative (catégorielle)", "Fréquence d'accès au crédit (Une fois, Deux fois, Trois fois ou plus)."),
    ("Q10-Montant total du dernier crédit contracté", "Quantitative (continue)", "Montant du crédit en gourdes."),
    ("Q11-Temps écoulé entre la demande et l'approbation du crédit.", "Quantitative (continue)", "Temps d'approbation en jours."),
    ("Q13- Avez-vous investi au moins 80 % du prêt reçu dans l'agriculture?", "Qualitative (catégorielle)", "Utilisation du crédit (Oui/Non)."),
    ("Q16- Avez-vous constaté une augmentation de vos revenus agricoles depuis l'obtention du crédit?", "Qualitative (catégorielle)", "Impact économique (Oui/Non)."),
    ("Q19-Avez-vous remboursé intégralement votre dernier crédit agricole?", "Qualitative (catégorielle)", "Remboursement du crédit (Oui/Non)."),
    ("Q23- Combien de transactions commerciales avez-vous effectuées avec des acheteurs externes au cours des 12 derniers mois?", "Qualitative (catégorielle)", "Transactions commerciales (Aucune, 1 à 2, 3 et plus)."),
    ("Q25- Avez-vous bénéficié d'un appui des comités mis en place pour la gestion des crédits et des risques financiers?", "Qualitative (catégorielle)", "Soutien institutionnel (Oui/Non)."),
]

# Remplir le tableau avec les données
for var_name, var_type, var_comment in variables:
    row_cells = table.add_row().cells
    row_cells[0].text = var_name
    row_cells[1].text = var_type
    row_cells[2].text = var_comment

# Ajouter un sous-titre pour l'explication des variables
doc.add_heading('Explication des variables', level=2)

# Ajouter un paragraphe d'explication
doc.add_paragraph(
    "1. **Âge (Q2)** : Cette variable quantitative permet d'analyser la distribution des âges des répondants et de voir si l'âge influence l'accès au crédit ou son utilisation.\n"
    "2. **Sexe (Q3)** : Cette variable catégorielle permet d'étudier les différences entre hommes et femmes en termes d'accès au crédit, d'utilisation et d'impact.\n"
    "3. **Taille du ménage (Q5)** : Cette variable quantitative donne une idée de la structure des ménages et peut être corrélée avec les besoins en crédit.\n"
    "4. **Accès au crédit (Q8)** : Cette variable catégorielle permet de savoir si les répondants ont déjà contracté un crédit agricole.\n"
    "5. **Fréquence d'accès au crédit (Q9)** : Cette variable catégorielle indique combien de fois les répondants ont obtenu un crédit agricole au cours des trois dernières années.\n"
    "6. **Montant du crédit (Q10)** : Cette variable quantitative mesure le montant du dernier crédit contracté, ce qui permet d'analyser l'ampleur des investissements.\n"
    "7. **Temps d'approbation (Q11)** : Cette variable quantitative mesure le temps écoulé entre la demande et l'approbation du crédit, ce qui peut influencer la satisfaction des agriculteurs.\n"
    "8. **Utilisation du crédit (Q13)** : Cette variable catégorielle permet de savoir si les agriculteurs ont investi au moins 80 % du prêt dans l'agriculture.\n"
    "9. **Augmentation des revenus (Q16)** : Cette variable catégorielle permet de mesurer l'impact économique du crédit sur les revenus des agriculteurs.\n"
    "10. **Remboursement du crédit (Q19)** : Cette variable catégorielle permet de savoir si les agriculteurs ont remboursé intégralement leur dernier crédit.\n"
    "11. **Transactions commerciales (Q23)** : Cette variable catégorielle mesure le nombre de transactions commerciales avec des acheteurs externes, ce qui reflète l'impact géographique du crédit.\n"
    "12. **Appui des comités (Q25)** : Cette variable catégorielle permet de savoir si les agriculteurs ont bénéficié d'un appui institutionnel pour la gestion des crédits."
)

# Ajouter un sous-titre pour la justification des choix
doc.add_heading('Justification des choix', level=2)

# Ajouter un paragraphe de justification
doc.add_paragraph(
    "Ces variables ont été sélectionnées car elles couvrent les trois domaines mentionnés dans le document Word :\n"
    "- **Socio-économique** : Âge, sexe, taille du ménage, accès au crédit, montant du crédit, utilisation du crédit, augmentation des revenus, remboursement du crédit.\n"
    "- **Politique** : Temps d'approbation, appui des comités.\n"
    "- **Géographique** : Transactions commerciales avec des acheteurs externes.\n\n"
    "Ces variables permettront de répondre aux objectifs du devoir en fournissant une analyse complète de l'impact des crédits agricoles sur les agriculteurs de Thomassique."
)

# Enregistrer le document
doc.save('Etape1_Selection_Variables.docx')

